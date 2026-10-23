"""
Генерация xlsx-отчёта «ушли в ноль» для Telegram (docs/push_message_about_zero.md).
"""

from __future__ import annotations

import logging
import os
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.zero_sold import ZeroSoldHit

logger = logging.getLogger(__name__)

_REASON_LABEL = {
    "zero": "явный ноль",
    "missing": "нет в файле",
}


def generate_zero_sold_xlsx(
    hits: Sequence[ZeroSoldHit],
    output_path: str,
) -> str:
    """
    Пишет таблицу ушедших в ноль позиций в ``.xlsx``.

    Колонки: Группа, Бренд, Наименование, Было (шт.), Стало, Причина.

    :param hits: Результат :func:`services.zero_sold.find_zero_sold`.
    :type hits: Sequence[ZeroSoldHit]
    :param output_path: Полный путь файла для записи.
    :type output_path: str
    :returns: Тот же ``output_path``.
    :rtype: str
    :side effects: Создаёт/перезаписывает файл на диске.
    """
    directory = os.path.dirname(os.path.abspath(output_path))
    if directory:
        os.makedirs(directory, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ушли в 0"

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    headers = [
        "Группа",
        "Бренд",
        "Наименование",
        "Было (шт.)",
        "Стало",
        "Причина",
    ]
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    for row_idx, hit in enumerate(hits, start=2):
        became = "0" if hit.reason == "zero" else "—"
        reason = _REASON_LABEL.get(hit.reason, hit.reason)
        row_vals = [
            hit.group,
            hit.brand,
            hit.name,
            hit.qty_was,
            became,
            reason,
        ]
        for col, value in enumerate(row_vals, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 3))

    widths = [28, 16, 48, 12, 10, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    wb.save(output_path)
    logger.info("Zero-sold xlsx сохранён: %s (%d поз.)", output_path, len(hits))
    return output_path
