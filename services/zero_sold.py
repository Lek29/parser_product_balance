"""
Zero-sold: снимок недели → diff с прошлой → xlsx в личку (K1, A+B).

См. ``docs/push_message_about_zero.md``.
"""

from __future__ import annotations

import json
import logging
import os
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Literal, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from parser.excel_parser import normalize_top_category_label
from services.calendar_week import (
    now_moscow,
    previous_week_monday,
    today_moscow,
    week_monday,
)
from storage.data_store import atomic_write_json

logger = logging.getLogger(__name__)

Reason = Literal["zero", "missing"]
_REASON_LABEL = {"zero": "явный ноль", "missing": "нет в файле"}


@dataclass(frozen=True)
class ZeroSoldHit:
    """Позиция: на прошлой неделе qty > 0, сейчас 0 или строки нет."""

    group: str
    brand: str
    name: str
    qty_was: int
    reason: Reason


def make_position_key(item: dict) -> str:
    """
    Ключ K1: группа + бренд + название (casefold/strip).

    :param item: Позиция ``parse_stock`` (``top_category``) или снимок (``group``).
    :type item: dict
    :returns: Стабильный ключ.
    :rtype: str
    """
    raw_group = item.get("top_category")
    if raw_group is None:
        raw_group = item.get("group", "")
    group = normalize_top_category_label(str(raw_group or ""))
    brand = str(item.get("brand") or "").strip().casefold()
    name = str(item.get("name") or "").strip().casefold()
    return f"{group.casefold()}|{brand}|{name}"


def _to_map(rows: list[dict]) -> dict[str, dict]:
    """key → {group, brand, name, quantity}; дубликаты — max qty."""
    out: dict[str, dict] = {}
    for raw in rows:
        qty_raw = raw.get("quantity")
        if qty_raw is None:
            continue
        try:
            qty = int(qty_raw)
        except (TypeError, ValueError):
            continue
        # снимок уже с key/group; parse_stock — top_category
        if raw.get("key") is not None and "top_category" not in raw:
            key = str(raw["key"])
            group = str(raw.get("group") or "")
            brand = str(raw.get("brand") or "")
            name = str(raw.get("name") or "")
        else:
            key = make_position_key(raw)
            group = normalize_top_category_label(str(raw.get("top_category", "") or ""))
            brand = str(raw.get("brand") or "").strip()
            name = str(raw.get("name") or "").strip()
        if not name:
            continue
        entry = {"group": group, "brand": brand, "name": name, "quantity": qty}
        prev = out.get(key)
        if prev is None or qty > int(prev["quantity"]):
            out[key] = entry
    return out


def find_zero_sold(prev_positions: list[dict], curr_items: list[dict]) -> list[ZeroSoldHit]:
    """
    Было qty > 0 → сейчас 0 или позиции нет.

    :param prev_positions: Позиции снимка прошлой недели.
    :type prev_positions: list[dict]
    :param curr_items: Текущий ``parse_stock``.
    :type curr_items: list[dict]
    :returns: Hit'ы, сортировка по группе/имени.
    :rtype: list[ZeroSoldHit]
    """
    prev_map = _to_map(prev_positions)
    curr_map = _to_map(curr_items)
    hits: list[ZeroSoldHit] = []
    for key, prev in prev_map.items():
        qty_was = int(prev["quantity"])
        if qty_was <= 0:
            continue
        if key not in curr_map:
            hits.append(
                ZeroSoldHit(
                    group=prev["group"],
                    brand=prev["brand"],
                    name=prev["name"],
                    qty_was=qty_was,
                    reason="missing",
                )
            )
        elif int(curr_map[key]["quantity"]) == 0:
            hits.append(
                ZeroSoldHit(
                    group=prev["group"],
                    brand=prev["brand"],
                    name=prev["name"],
                    qty_was=qty_was,
                    reason="zero",
                )
            )
    hits.sort(key=lambda h: (h.group.casefold(), h.name.casefold(), h.brand.casefold()))
    return hits


def _snapshot_path(week_id: str) -> str:
    return os.path.join(config.WEEKLY_SNAPSHOTS_DIR, f"{week_id.strip()}.json")


def _save_week_snapshot(week_id: str, items: list[dict]) -> None:
    """Пишет снимок недели из parse_stock items."""
    by_key = _to_map(items)
    positions = [
        {
            "key": k,
            "group": v["group"],
            "brand": v["brand"],
            "name": v["name"],
            "quantity": v["quantity"],
        }
        for k, v in by_key.items()
    ]
    os.makedirs(config.WEEKLY_SNAPSHOTS_DIR, exist_ok=True)
    atomic_write_json(
        _snapshot_path(week_id),
        {
            "week_id": week_id,
            "saved_at": now_moscow().isoformat(),
            "positions": positions,
        },
        indent=2,
        ensure_ascii=False,
    )
    logger.info("zero_sold snapshot week_id=%s n=%d", week_id, len(positions))


def _load_week_snapshot(week_id: str) -> Optional[dict[str, Any]]:
    path = _snapshot_path(week_id)
    if not os.path.isfile(path):
        return None
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except (OSError, json.JSONDecodeError) as e:
        logger.error("zero_sold snapshot read fail %s: %s", path, e)
        return None
    return data if isinstance(data, dict) else None


def _load_state() -> dict[str, Any]:
    path = config.ZERO_SOLD_STATE_JSON_PATH
    if not os.path.isfile(path):
        return {}
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except (OSError, json.JSONDecodeError):
        return {}
    return data if isinstance(data, dict) else {}


def _save_state(week_from: str, week_to: str) -> None:
    atomic_write_json(
        config.ZERO_SOLD_STATE_JSON_PATH,
        {
            "last_compared_from": week_from,
            "last_compared_to": week_to,
            "pushed_at": datetime.now().isoformat(),
        },
        indent=2,
        ensure_ascii=False,
    )


def _write_xlsx(hits: list[ZeroSoldHit], path: str) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ушли в 0"
    fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    font = Font(bold=True, color="FFFFFF", size=11)
    headers = ["Группа", "Бренд", "Наименование", "Было (шт.)", "Стало", "Причина"]
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_idx, hit in enumerate(hits, start=2):
        vals = [
            hit.group,
            hit.brand,
            hit.name,
            hit.qty_was,
            "0" if hit.reason == "zero" else "—",
            _REASON_LABEL[hit.reason],
        ]
        for col, value in enumerate(vals, start=1):
            ws.cell(row=row_idx, column=col, value=value)
    for i, w in enumerate([28, 16, 48, 12, 10, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)


def run_zero_sold_check_after_parse(all_items: list[dict]) -> None:
    """
    Хук после полного парса: снимок → diff → xlsx в ``STOCK_LIMIT_ALERT_CHAT_ID``.

    :param all_items: Результат ``parse_stock``.
    :type all_items: list[dict]
    :returns: None
    :side effects: snapshots, state, temp xlsx, Telegram document.
    """
    try:
        _run(all_items)
    except Exception:
        logger.error("zero_sold fail", exc_info=True)


def _run(all_items: list[dict]) -> None:
    from bot.telegram_bridge import send_zero_sold_xlsx

    today = today_moscow()
    week_to = week_monday(today).isoformat()
    week_from = previous_week_monday(today).isoformat()

    _save_week_snapshot(week_to, all_items)

    state = _load_state()
    if state.get("last_compared_from") == week_from and state.get("last_compared_to") == week_to:
        logger.info("zero_sold skip already_compared %s→%s", week_from, week_to)
        return

    prev = _load_week_snapshot(week_from)
    if prev is None:
        logger.info("zero_sold skip no_prev week_from=%s", week_from)
        return

    prev_positions = prev.get("positions") or []
    if not isinstance(prev_positions, list):
        prev_positions = []

    hits = find_zero_sold(prev_positions, all_items)
    logger.info("zero_sold compare %s→%s n_hits=%d", week_from, week_to, len(hits))

    if not hits:
        _save_state(week_from, week_to)
        return

    filename = f"ушли_в_ноль_{today.isoformat()}.xlsx"
    path = os.path.join(config.TEMP_DIR, filename)
    _write_xlsx(hits, path)
    caption = f"Ушли в 0 относительно прошлой недели ({week_from} → {week_to}): {len(hits)} поз."
    sent = send_zero_sold_xlsx(path, caption=caption, filename=filename)
    if sent:
        _save_state(week_from, week_to)
    try:
        if os.path.isfile(path):
            os.remove(path)
    except OSError:
        pass
