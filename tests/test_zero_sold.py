"""
Тесты zero-sold: ключ K1, diff A+B, отчёт xlsx (без Telegram).
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from services.zero_sold import (
    ZeroSoldHit,
    find_zero_sold,
    make_position_key,
)
from services.zero_sold_report import generate_zero_sold_xlsx


def _item(
    group: str,
    brand: str,
    name: str,
    qty: int,
) -> dict:
    """Минимальная позиция parse_stock."""
    return {
        "top_category": group,
        "subcategory": "",
        "brand": brand,
        "name": name,
        "quantity": qty,
        "price": 1,
    }


def test_make_position_key_k1_casefold() -> None:
    """K1: группа+бренд+имя, без учёта регистра имени/бренда."""
    a = _item("Аксессуары", "Atemi", "Зажим X", 2)
    b = _item("Аксессуары", "ATEMI", "зажим x", 1)
    assert make_position_key(a) == make_position_key(b)


def test_find_zero_sold_explicit_zero() -> None:
    """Было >0, стало 0 → reason=zero."""
    prev = [
        {
            "key": make_position_key(_item("G", "B", "N", 2)),
            "group": "G",
            "brand": "B",
            "name": "N",
            "quantity": 2,
        }
    ]
    curr = [_item("G", "B", "N", 0)]
    hits = find_zero_sold(prev, curr)
    assert len(hits) == 1
    assert hits[0].reason == "zero"
    assert hits[0].qty_was == 2


def test_find_zero_sold_missing_row() -> None:
    """Было >0, строки нет → reason=missing."""
    prev = [
        {
            "key": make_position_key(_item("G", "B", "Gone", 3)),
            "group": "G",
            "brand": "B",
            "name": "Gone",
            "quantity": 3,
        }
    ]
    hits = find_zero_sold(prev, [_item("G", "B", "Other", 1)])
    assert len(hits) == 1
    assert hits[0].reason == "missing"
    assert hits[0].name == "Gone"


def test_find_zero_sold_no_false_positives() -> None:
    """0→0, новая позиция, стабильный qty — без hit."""
    prev = [
        {
            "key": make_position_key(_item("G", "B", "Zero", 0)),
            "group": "G",
            "brand": "B",
            "name": "Zero",
            "quantity": 0,
        },
        {
            "key": make_position_key(_item("G", "B", "Keep", 5)),
            "group": "G",
            "brand": "B",
            "name": "Keep",
            "quantity": 5,
        },
    ]
    curr = [
        _item("G", "B", "Keep", 5),
        _item("G", "B", "New", 1),
        _item("G", "B", "Zero", 0),
    ]
    assert find_zero_sold(prev, curr) == []


def test_find_zero_sold_empty_prev() -> None:
    """Нет prev позиций — пустой список."""
    assert find_zero_sold([], [_item("G", "B", "N", 0)]) == []


def test_generate_zero_sold_xlsx(tmp_path: Path) -> None:
    """Файл создаётся с ожидаемыми заголовками и строкой."""
    hits = [
        ZeroSoldHit(
            key="k",
            group="Группа",
            brand="Brand",
            name="Товар",
            qty_was=4,
            reason="missing",
        )
    ]
    out = tmp_path / "ушли_в_ноль_2026-10-23.xlsx"
    generate_zero_sold_xlsx(hits, str(out))
    assert out.is_file()
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert ws.cell(1, 1).value == "Группа"
    assert ws.cell(2, 3).value == "Товар"
    assert ws.cell(2, 4).value == 4
    assert ws.cell(2, 6).value == "нет в файле"
