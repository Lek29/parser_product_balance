"""
Тесты календарной недели МСК, свежести IMAP-skip и кэша парсинга (P1).
"""

from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
import pytest
from openpyxl.styles import Alignment

import config
from parser.excel_parser import HEADER_ROWS


def test_week_monday_and_same_week() -> None:
    """Пн–Вс: вторник и воскресенье одной недели; понедельник следующей — нет."""
    from services.calendar_week import same_calendar_week, week_monday

    tue = date(2026, 10, 20)  # вт
    sun = date(2026, 10, 25)  # вс
    next_mon = date(2026, 10, 26)
    assert week_monday(tue) == date(2026, 10, 19)
    assert same_calendar_week(tue, sun) is True
    assert same_calendar_week(tue, next_mon) is False


def test_parse_meta_date_iso_and_date() -> None:
    """Разбор YYYY-MM-DD и ISO datetime (naive = МСК)."""
    from services.calendar_week import parse_meta_date

    assert parse_meta_date("2026-10-22") == date(2026, 10, 22)
    assert parse_meta_date("2026-10-22T15:30:00") == date(2026, 10, 22)
    assert parse_meta_date("") is None
    assert parse_meta_date(None) is None


def test_local_snapshot_fresh_same_week(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """current.xlsx + last_check_date в той же неделе МСК → fresh."""
    xlsx = tmp_path / "current.xlsx"
    xlsx.write_bytes(b"PK\x03\x04fake")
    meta = tmp_path / "meta.json"
    meta.write_text(
        json.dumps({"last_check_date": "2026-10-21", "updated_at": "2026-10-21T10:00:00"}),
        encoding="utf-8",
    )
    monkeypatch.setattr(config, "CURRENT_XLSX_PATH", str(xlsx))
    monkeypatch.setattr(config, "META_JSON_PATH", str(meta))

    from storage.data_store import is_local_mail_snapshot_fresh_this_week

    ref = datetime(2026, 10, 23, 12, 0, tzinfo=ZoneInfo("Europe/Moscow"))
    assert is_local_mail_snapshot_fresh_this_week(reference=ref) is True

    ref_next = datetime(2026, 10, 26, 9, 0, tzinfo=ZoneInfo("Europe/Moscow"))
    assert is_local_mail_snapshot_fresh_this_week(reference=ref_next) is False


def test_data_json_fingerprint_cache(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """save_items пишет отпечаток; совпадение с current.xlsx → cache hit."""
    xlsx = tmp_path / "current.xlsx"
    _write_minimal_stock_xlsx(xlsx)
    data_path = tmp_path / "data.json"
    monkeypatch.setattr(config, "CURRENT_XLSX_PATH", str(xlsx))
    monkeypatch.setattr(config, "DATA_JSON_PATH", str(data_path))

    from storage.data_store import (
        current_xlsx_fingerprint,
        data_json_matches_current_xlsx,
        save_items,
    )

    assert data_json_matches_current_xlsx() is False
    fp = current_xlsx_fingerprint()
    assert fp is not None
    save_items(
        [{"top_category": "G", "name": "n", "quantity": 1, "price": 1}],
        source_fingerprint=fp,
    )
    assert data_json_matches_current_xlsx() is True

    # Смена файла — кэш невалиден
    xlsx.write_bytes(xlsx.read_bytes() + b"x")
    assert data_json_matches_current_xlsx() is False


def test_run_parse_uses_cache_skips_parse_stock(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """При валидном кэше parse_stock не вызывается."""
    xlsx = tmp_path / "current.xlsx"
    _write_minimal_stock_xlsx(xlsx)
    data_path = tmp_path / "data.json"
    filtered = tmp_path / "filtered.xlsx"
    filtered.write_bytes(b"PK\x03\x04cached")
    monkeypatch.setattr(config, "CURRENT_XLSX_PATH", str(xlsx))
    monkeypatch.setattr(config, "DATA_JSON_PATH", str(data_path))
    monkeypatch.setattr(config, "FILTERED_XLSX_PATH", str(filtered))

    from storage.data_store import current_xlsx_fingerprint, save_items

    items = [
        {
            "top_category": "G",
            "subcategory": "",
            "brand": "",
            "name": "n",
            "quantity": 1,
            "price": 10,
        }
    ]
    save_items(items, source_fingerprint=current_xlsx_fingerprint())

    parse_calls: list[int] = []

    def boom_parse(*_a, **_k):
        parse_calls.append(1)
        raise AssertionError("parse_stock не должен вызываться при cache hit")

    monkeypatch.setattr("services.stock_pipeline.parse_stock", boom_parse)
    monkeypatch.setattr("services.stock_pipeline.list_all_limits", dict)

    from services.stock_pipeline import run_parse_and_save_filtered

    all_n, filt_n = run_parse_and_save_filtered()
    assert parse_calls == []
    assert all_n == 1
    assert filt_n == 1


def test_filtered_xlsx_ready_for_send(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Готовый filtered + валидный data.json → ready."""
    xlsx = tmp_path / "current.xlsx"
    _write_minimal_stock_xlsx(xlsx)
    data_path = tmp_path / "data.json"
    filtered = tmp_path / "filtered.xlsx"
    filtered.write_bytes(b"PK\x03\x04ok")
    monkeypatch.setattr(config, "CURRENT_XLSX_PATH", str(xlsx))
    monkeypatch.setattr(config, "DATA_JSON_PATH", str(data_path))
    monkeypatch.setattr(config, "FILTERED_XLSX_PATH", str(filtered))

    from storage.data_store import (
        current_xlsx_fingerprint,
        filtered_xlsx_ready_for_send,
        save_items,
    )

    assert filtered_xlsx_ready_for_send() is False
    save_items(
        [{"top_category": "G", "name": "n", "quantity": 1, "price": 1}],
        source_fingerprint=current_xlsx_fingerprint(),
    )
    assert filtered_xlsx_ready_for_send() is True


def _write_minimal_stock_xlsx(path: Path) -> None:
    """Минимальный xlsx с одной позицией (indent 10)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    r_cat = HEADER_ROWS + 1
    cat_cell = ws.cell(row=r_cat, column=1, value="Группа")
    cat_cell.alignment = Alignment(indent=4)
    r_item = r_cat + 1
    item_cell = ws.cell(row=r_item, column=1, value="Товар 1")
    item_cell.alignment = Alignment(indent=10)
    ws.cell(row=r_item, column=4, value=2)
    ws.cell(row=r_item, column=5, value=50)
    wb.save(path)
