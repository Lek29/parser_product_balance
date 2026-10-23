"""
Минимальные unit-тесты без сети и Telegram (см. docs/refactoring-plan.md §2.3).
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Alignment

from parser.excel_parser import (
    HEADER_ROWS,
    TOP_CATEGORY_FALLBACK_LABEL,
    filter_by_group_thresholds,
    find_group_limit_violations,
    generate_filtered_xlsx,
    parse_stock,
)


def test_filter_by_group_thresholds_basic() -> None:
    """
    Частичный фильтр: только группы из словаря; порог включительно.
    """
    items = [
        {"top_category": "Аксессуары", "quantity": 2, "name": "a"},
        {"top_category": "Мячи", "quantity": 10, "name": "b"},
        {"top_category": "Аксессуары", "quantity": 7, "name": "c"},
    ]
    out = filter_by_group_thresholds(
        items,
        {"Аксессуары": 5, "Мячи": 15},
    )
    assert len(out) == 2
    assert {o["top_category"] for o in out} == {"Аксессуары", "Мячи"}


def test_filter_by_group_thresholds_empty_rules() -> None:
    """Пустой словарь порогов — пустой результат."""
    items = [{"top_category": "X", "quantity": 1}]
    assert filter_by_group_thresholds(items, {}) == []


def test_filter_by_group_thresholds_fallback_group() -> None:
    """Пустой top_category нормализуется в TOP_CATEGORY_FALLBACK_LABEL."""
    items = [{"top_category": "", "quantity": 1}]
    out = filter_by_group_thresholds(items, {TOP_CATEGORY_FALLBACK_LABEL: 5})
    assert len(out) == 1


def test_find_group_limit_violations_basic() -> None:
    """Группировка нарушений по группе; только позиции с quantity <= лимита."""
    items = [
        {"top_category": "Велосипеды", "name": "A", "quantity": 2},
        {"top_category": "Велосипеды", "name": "B", "quantity": 10},
        {"top_category": "Мячи", "name": "C", "quantity": 1},
    ]
    out = find_group_limit_violations(items, {"Велосипеды": 5})
    assert len(out) == 1
    assert out[0]["group"] == "Велосипеды"
    assert out[0]["limit"] == 5
    assert out[0]["positions"] == [{"name": "A", "quantity": 2}]


def test_find_group_limit_violations_empty_limits() -> None:
    """Пустой словарь лимитов — пустой список срабатываний."""
    items = [{"top_category": "X", "name": "n", "quantity": 0}]
    assert find_group_limit_violations(items, {}) == []


def test_find_group_limit_violations_no_hits() -> None:
    """Ни одна позиция не проходит порог — пустой результат."""
    items = [{"top_category": "Велосипеды", "name": "B", "quantity": 10}]
    assert find_group_limit_violations(items, {"Велосипеды": 3}) == []


def test_parse_group_limit_callback() -> None:
    """Разбор callback лимитов по группам (пресеты / свой / снять)."""
    from bot.keyboards import parse_group_limit_callback

    assert parse_group_limit_callback("glt:2:10") == (2, 10)
    assert parse_group_limit_callback("glt:0:x") == (0, "custom")
    assert parse_group_limit_callback("glt:1:c") == (1, "clear")
    assert parse_group_limit_callback("bad") is None


def test_is_known_product_group_reply_text(monkeypatch: pytest.MonkeyPatch) -> None:
    """Текст совпадает с группой из list_product_groups (патч без xlsx)."""
    import bot.partial_filter_state as pfs

    monkeypatch.setattr(
        "parser.excel_parser.list_product_groups",
        lambda xlsx_path=None: ["Альпинизм", "VELO"],
    )
    assert pfs.is_known_product_group_reply_text("Альпинизм") is True
    assert pfs.is_known_product_group_reply_text("velo") is True
    assert pfs.is_known_product_group_reply_text("NoSuchGroup") is False


def test_group_index_by_reply_label_casefold_latin() -> None:
    """Латинское имя группы совпадает с reply-текстом без учёта регистра."""
    from types import SimpleNamespace

    from bot.partial_filter_state import (
        UD_PARTIAL_GROUPS_ORDER,
        group_index_by_reply_label,
    )

    ctx = SimpleNamespace(
        user_data={UD_PARTIAL_GROUPS_ORDER: ["VELO", "Альпинизм"]},
    )
    assert group_index_by_reply_label(ctx, "velo") == 0
    assert group_index_by_reply_label(ctx, "Альпинизм") == 1


def test_normalize_reply_button_text_strips_variation_selector() -> None:
    """Emoji «✅» с U+FE0F сравнивается с константой без него."""
    from bot.keyboards import (
        PARTIAL_REPLY_APPLY_TEXT,
        normalize_reply_button_text,
    )

    with_vs16 = "✅\ufe0f Применить"
    assert normalize_reply_button_text(with_vs16) == normalize_reply_button_text(
        PARTIAL_REPLY_APPLY_TEXT,
    )


def test_normalize_reply_button_text_nbsp_like_apply() -> None:
    """Неразрывный пробел между emoji и словом — как у обычного пробела."""
    from bot.keyboards import (
        PARTIAL_REPLY_APPLY_TEXT,
        normalize_reply_button_text,
    )

    nbsp = "✅\u00a0Применить"
    assert normalize_reply_button_text(nbsp) == normalize_reply_button_text(
        PARTIAL_REPLY_APPLY_TEXT,
    )


def test_format_group_limit_alerts_html_basic_and_escape() -> None:
    """HTML-алерт: заголовок группы, список, экранирование спецсимволов."""
    from services.group_limit_alert_text import format_group_limit_alerts_html

    violations = [
        {
            "group": "Вело <test>",
            "limit": 5,
            "positions": [
                {"name": 'SKU "A"', "quantity": 2},
                {"name": "B", "quantity": 5},
            ],
        },
    ]
    text = format_group_limit_alerts_html(violations)
    assert text is not None
    assert "&lt;" in text and "&quot;" in text
    assert "Вело" in text
    assert "<code>5</code>" in text


def test_format_group_limit_alerts_html_empty() -> None:
    """Пустой вход — None."""
    from services.group_limit_alert_text import format_group_limit_alerts_html

    assert format_group_limit_alerts_html([]) is None
    assert (
        format_group_limit_alerts_html(
            [{"group": "X", "limit": 1, "positions": []}],
        )
        is None
    )


def test_format_group_limit_alerts_html_truncates_omitted_line() -> None:
    """При большом числе позиций — строка «и ещё N поз.»."""
    from services.group_limit_alert_text import format_group_limit_alerts_html

    positions = [{"name": f"P{i}", "quantity": 1} for i in range(80)]
    violations = [{"group": "G", "limit": 3, "positions": positions}]
    text = format_group_limit_alerts_html(
        violations,
        max_message_chars=450,
    )
    assert text is not None
    assert "и ещё" in text


def test_send_group_limit_alert_html_no_chat_no_crash(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Без chat_id и без polling — отправка алерта не падает."""
    import config

    monkeypatch.setattr(config, "STOCK_LIMIT_ALERT_CHAT_ID", None)
    from bot.telegram_bridge import send_group_limit_alert_html

    send_group_limit_alert_html("<b>test</b>")


def test_check_update_available_no_meta(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Нет meta.json — обновление «нужно»."""
    import config

    monkeypatch.setattr(config, "META_JSON_PATH", str(tmp_path / "missing.json"))
    from storage.data_store import check_update_available

    assert check_update_available("Mon, 1 Jan 2026 00:00:00 +0000") is True


def test_check_update_available_same_date_string(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Совпадение строки даты — без обновления."""
    import config

    meta_path = tmp_path / "meta.json"
    date_str = "Mon, 16 Mar 2026 09:00:00 +0800"
    meta_path.write_text(json.dumps({"email_date": date_str}), encoding="utf-8")
    monkeypatch.setattr(config, "META_JSON_PATH", str(meta_path))

    from storage.data_store import check_update_available

    assert check_update_available(date_str) is False


def test_load_items_corrupt_json(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Повреждённый data.json — пустой список, без исключения (QA-002)."""
    import config

    bad = tmp_path / "data.json"
    bad.write_text("{not json", encoding="utf-8")
    monkeypatch.setattr(config, "DATA_JSON_PATH", str(bad))

    from storage.data_store import load_items

    assert load_items() == []


def test_atomic_write_json_roundtrip(tmp_path: Path) -> None:
    """Атомарная запись читается обратно."""
    p = tmp_path / "t.json"
    from storage.data_store import atomic_write_json

    atomic_write_json(str(p), {"a": 1}, indent=2)
    assert json.loads(p.read_text(encoding="utf-8")) == {"a": 1}


def test_generate_filtered_xlsx_empty_top_category_uses_fallback_label(
    tmp_path: Path,
) -> None:
    """
    Пустой ``top_category`` в выгрузке даёт заголовок группы как в частичном фильтре.
    """
    out = tmp_path / "out.xlsx"
    items = [
        {
            "top_category": "",
            "subcategory": "",
            "brand": "B",
            "name": "SKU",
            "quantity": 1,
            "price": 100,
        },
    ]
    path = generate_filtered_xlsx(items, output_path=str(out))
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    header = ws.cell(row=2, column=1).value
    assert header is not None
    assert TOP_CATEGORY_FALLBACK_LABEL in str(header)


def test_authorize_idempotent(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Повторный authorize для того же chat_id не дублирует запись в auth.json."""
    import config
    import storage.data_store as ds

    auth_path = tmp_path / "auth.json"
    monkeypatch.setattr(config, "AUTH_JSON_PATH", str(auth_path))
    ds._auth_ids_cache = None
    ds._auth_cache_mtime = None

    from storage.data_store import authorize

    authorize(999001)
    authorize(999001)

    data = json.loads(auth_path.read_text(encoding="utf-8"))
    assert data["authorized"] == [999001]
    assert data["authorized"].count(999001) == 1


def test_read_meta_json_unifies_status_fields(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Один объект meta даёт согласованные поля для get_last_email_date_str и статуса."""
    import config

    meta_path = tmp_path / "meta.json"
    payload = {
        "email_date": "Mon, 1 Jan 2026 12:00:00 +0000",
        "updated_at": "2026-01-01T15:00:00",
        "last_check_date": "2026-01-01",
    }
    meta_path.write_text(json.dumps(payload), encoding="utf-8")
    monkeypatch.setattr(config, "META_JSON_PATH", str(meta_path))

    from storage.data_store import (
        get_last_email_date_str,
        get_status_text,
        read_meta_json,
    )

    m = read_meta_json()
    assert m is not None
    assert m.get("email_date") == payload["email_date"]
    assert get_last_email_date_str() == payload["email_date"]
    assert "01.01.2026" in get_status_text()


def test_group_stock_limits_roundtrip_and_normalization(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """
    Лимиты по группам сохраняются атомарно, нормализуют имя группы и читаются обратно.
    """
    import config

    limits_path = tmp_path / "group_stock_limits.json"
    monkeypatch.setattr(
        config,
        "GROUP_STOCK_LIMITS_JSON_PATH",
        str(limits_path),
    )

    from storage.data_store import get_limit, list_all_limits, set_limit

    set_limit(" Велосипеды ", 3)

    payload = json.loads(limits_path.read_text(encoding="utf-8"))
    assert payload["count"] == 1
    assert payload["limits"] == {"Велосипеды": 3}
    assert get_limit("Велосипеды") == 3
    assert list_all_limits() == {"Велосипеды": 3}


def test_clear_limit_removes_group_from_registry(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """
    clear_limit удаляет ключ из словаря лимитов и снимает группу с надзора.
    """
    import config

    limits_path = tmp_path / "group_stock_limits.json"
    monkeypatch.setattr(
        config,
        "GROUP_STOCK_LIMITS_JSON_PATH",
        str(limits_path),
    )

    from storage.data_store import (
        clear_limit,
        get_limit,
        list_all_limits,
        set_limit,
    )

    set_limit("Велосипеды", 3)
    set_limit("", 0)
    clear_limit(" Велосипеды ")

    assert get_limit("Велосипеды") is None
    assert list_all_limits() == {TOP_CATEGORY_FALLBACK_LABEL: 0}


def test_group_stock_limits_corrupt_json_returns_empty(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """
    Повреждённый group_stock_limits.json
    не валит рантайм и даёт пустой словарь.
    """
    import config

    limits_path = tmp_path / "group_stock_limits.json"
    limits_path.write_text("{not json", encoding="utf-8")
    monkeypatch.setattr(
        config,
        "GROUP_STOCK_LIMITS_JSON_PATH",
        str(limits_path),
    )

    from storage.data_store import get_limit, list_all_limits

    assert list_all_limits() == {}
    assert get_limit("Велосипеды") is None


def test_parse_stock_read_only_preserves_indent_hierarchy(tmp_path: Path) -> None:
    """read_only parse: категории (indent 4) и SKU (indent 10) связываются верно."""
    path = tmp_path / "stock_indent.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    r = HEADER_ROWS + 1
    cat = ws.cell(row=r, column=1, value="Категория А")
    cat.alignment = Alignment(indent=4)
    r += 1
    item = ws.cell(row=r, column=1, value="Товар А1")
    item.alignment = Alignment(indent=10)
    ws.cell(row=r, column=4, value=3)
    ws.cell(row=r, column=5, value=100)
    r += 1
    cat2 = ws.cell(row=r, column=1, value="Категория Б")
    cat2.alignment = Alignment(indent=4)
    r += 1
    item2 = ws.cell(row=r, column=1, value="Товар Б1")
    item2.alignment = Alignment(indent=10)
    ws.cell(row=r, column=4, value=1)
    ws.cell(row=r, column=5, value=50)
    wb.save(path)

    items = parse_stock(str(path))
    assert len(items) == 2
    assert items[0]["top_category"] == "Категория А"
    assert items[0]["name"] == "Товар А1"
    assert items[0]["quantity"] == 3
    assert items[1]["top_category"] == "Категория Б"
    assert items[1]["name"] == "Товар Б1"


def test_parse_stock_always_reads_columns_a_through_e(tmp_path: Path) -> None:
    """
    Строки, где в файле заполнена только A (категория), не должны давать IndexError
    при обращении к D/E — iter_rows с max_col=5 отдаёт пустые ячейки.
    """
    path = tmp_path / "stock.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    r_cat = HEADER_ROWS + 1
    cat_cell = ws.cell(row=r_cat, column=1, value="Верх")
    cat_cell.alignment = Alignment(indent=4)
    r_item = r_cat + 1
    item_cell = ws.cell(row=r_item, column=1, value="Товар 1")
    item_cell.alignment = Alignment(indent=10)
    ws.cell(row=r_item, column=4, value=2)
    ws.cell(row=r_item, column=5, value=50)
    wb.save(path)

    items = parse_stock(str(path))
    assert len(items) == 1
    assert items[0]["name"] == "Товар 1"
    assert items[0]["quantity"] == 2


def test_run_parse_pipeline_single_telegram_send_for_multiple_groups(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """
    Этап 7: один прогон парсинга — одна отправка (HTML по всем группам).
    """
    items = [
        {"top_category": "Альфа", "name": "s1", "quantity": 1},
        {"top_category": "Бета", "name": "s2", "quantity": 2},
    ]
    sends: list[str] = []

    monkeypatch.setattr("services.stock_pipeline.parse_stock", lambda: items)
    monkeypatch.setattr(
        "services.stock_pipeline.data_json_matches_current_xlsx",
        lambda: False,
    )
    monkeypatch.setattr(
        "services.stock_pipeline.save_items",
        lambda _x, **_kw: None,
    )
    monkeypatch.setattr(
        "services.stock_pipeline.filter_by_threshold",
        lambda _x: [],
    )
    monkeypatch.setattr(
        "services.stock_pipeline.generate_filtered_xlsx",
        lambda _x: None,
    )
    monkeypatch.setattr(
        "services.zero_sold.run_zero_sold_check_after_parse",
        lambda _items: None,
    )
    monkeypatch.setattr(
        "services.stock_pipeline.list_all_limits",
        lambda: {"Альфа": 5, "Бета": 5},
    )
    monkeypatch.setattr(
        "services.stock_pipeline.send_group_limit_alert_html",
        lambda html: sends.append(html),
    )

    from services.stock_pipeline import run_parse_and_save_filtered

    all_n, filt_n = run_parse_and_save_filtered()
    assert all_n == 2
    assert filt_n == 0
    assert len(sends) == 1
    assert "Альфа" in sends[0] and "Бета" in sends[0]


def test_run_parse_pipeline_no_alert_when_no_limits(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """При пустом словаре лимитов ветка отправки не вызывается."""
    sends: list[str] = []

    monkeypatch.setattr(
        "services.stock_pipeline.parse_stock",
        lambda: [{"top_category": "X", "name": "n", "quantity": 1}],
    )
    monkeypatch.setattr(
        "services.stock_pipeline.data_json_matches_current_xlsx",
        lambda: False,
    )
    monkeypatch.setattr(
        "services.stock_pipeline.save_items",
        lambda _x, **_kw: None,
    )
    monkeypatch.setattr(
        "services.stock_pipeline.filter_by_threshold",
        lambda _x: [],
    )
    monkeypatch.setattr(
        "services.stock_pipeline.generate_filtered_xlsx",
        lambda _x: None,
    )
    monkeypatch.setattr(
        "services.zero_sold.run_zero_sold_check_after_parse",
        lambda _items: None,
    )
    monkeypatch.setattr("services.stock_pipeline.list_all_limits", dict)
    monkeypatch.setattr(
        "services.stock_pipeline.send_group_limit_alert_html",
        lambda html: sends.append(html),
    )

    from services.stock_pipeline import run_parse_and_save_filtered

    run_parse_and_save_filtered()
    assert sends == []


def _write_minimal_stock_xlsx(path: Path) -> None:
    """Один SKU (indent 10) для parse_stock."""
    wb = openpyxl.Workbook()
    ws = wb.active
    r_cat = HEADER_ROWS + 1
    cat_cell = ws.cell(row=r_cat, column=1, value="Группа")
    cat_cell.alignment = Alignment(indent=4)
    r_item = r_cat + 1
    item_cell = ws.cell(row=r_item, column=1, value="Товар 1")
    item_cell.alignment = Alignment(indent=10)
    ws.cell(row=r_item, column=4, value=2)
    ws.cell(row=r_item, column=5, value=50)
    wb.save(path)


def test_revision_skips_when_no_limits(monkeypatch: pytest.MonkeyPatch) -> None:
    """Этап 7а: пустой реестр лимитов — fetch и парсинг не вызываются."""
    fetch_calls: list[int] = []

    def boom() -> None:
        fetch_calls.append(1)
        raise AssertionError("fetch_and_save не должен вызываться")

    monkeypatch.setattr(
        "services.group_limit_revision.list_all_limits",
        dict,
    )
    monkeypatch.setattr(
        "services.group_limit_revision.fetch_and_save",
        boom,
    )
    from services.group_limit_revision import run_weekly_group_limit_revision

    run_weekly_group_limit_revision()
    assert fetch_calls == []


def test_local_snapshot_acceptable_last_check_is_today_moscow(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    from datetime import datetime
    from zoneinfo import ZoneInfo

    import config

    path = tmp_path / "current.xlsx"
    _write_minimal_stock_xlsx(path)
    monkeypatch.setattr(config, "CURRENT_XLSX_PATH", str(path))
    monkeypatch.setattr(
        "services.group_limit_revision.get_last_check_date",
        lambda: "2026-04-19",
    )
    ref = datetime(2026, 4, 19, 11, 0, tzinfo=ZoneInfo("Europe/Moscow"))
    from services.group_limit_revision import (
        local_snapshot_acceptable_for_limit_revision,
    )

    assert local_snapshot_acceptable_for_limit_revision(reference=ref) is True


def test_local_snapshot_acceptable_via_updated_at_meta(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    from datetime import datetime
    from zoneinfo import ZoneInfo

    import config

    path = tmp_path / "current.xlsx"
    _write_minimal_stock_xlsx(path)
    monkeypatch.setattr(config, "CURRENT_XLSX_PATH", str(path))
    monkeypatch.setattr(
        "services.group_limit_revision.get_last_check_date",
        lambda: "2026-04-01",
    )
    monkeypatch.setattr(
        "services.group_limit_revision.read_meta_json",
        lambda: {"updated_at": "2026-04-18T12:00:00"},
    )
    ref = datetime(2026, 4, 19, 12, 0, tzinfo=ZoneInfo("Europe/Moscow"))
    from services.group_limit_revision import (
        local_snapshot_acceptable_for_limit_revision,
    )

    assert local_snapshot_acceptable_for_limit_revision(reference=ref) is True


def test_run_weekly_revision_local_path_calls_send(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    items = [{"top_category": "А", "name": "x", "quantity": 1}]
    sent: list[list[dict]] = []

    monkeypatch.setattr(
        "services.group_limit_revision.list_all_limits",
        lambda: {"А": 5},
    )
    monkeypatch.setattr(
        "services.group_limit_revision.local_snapshot_acceptable_for_limit_revision",
        lambda reference=None: True,
    )
    monkeypatch.setattr(
        "services.group_limit_revision.parse_stock",
        lambda _p=None: items,
    )
    monkeypatch.setattr(
        "services.group_limit_revision.send_group_limit_alerts_for_items",
        lambda data, **_kw: sent.append(data),
    )

    from services.group_limit_revision import run_weekly_group_limit_revision

    run_weekly_group_limit_revision()
    assert len(sent) == 1 and sent[0] is items


def test_run_weekly_revision_updated_calls_full_pipeline(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from parser.mail_fetcher import MailFetchOutcome

    monkeypatch.setattr(
        "services.group_limit_revision.list_all_limits",
        lambda: {"А": 1},
    )
    monkeypatch.setattr(
        "services.group_limit_revision.local_snapshot_acceptable_for_limit_revision",
        lambda reference=None: False,
    )
    monkeypatch.setattr(
        "services.group_limit_revision.fetch_and_save",
        lambda: MailFetchOutcome.UPDATED,
    )
    called: list[str] = []

    def fake_run_parse() -> tuple[int, int]:
        called.append("run_parse")
        return (1, 0)

    monkeypatch.setattr(
        "services.group_limit_revision.run_parse_and_save_filtered",
        fake_run_parse,
    )

    from services.group_limit_revision import run_weekly_group_limit_revision

    run_weekly_group_limit_revision()
    assert called == ["run_parse"]


def test_send_group_limit_alerts_build_error_does_not_propagate(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Этап 8: сбой расчёта нарушений/HTML не ломает вызывающий код."""
    sends: list[str] = []

    def boom(*_a, **_kw):
        raise RuntimeError("build")

    monkeypatch.setattr("services.stock_pipeline.list_all_limits", lambda: {"A": 1})
    monkeypatch.setattr(
        "services.stock_pipeline.find_group_limit_violations",
        boom,
    )
    monkeypatch.setattr(
        "services.stock_pipeline.send_group_limit_alert_html",
        lambda h: sends.append(h),
    )
    from services.stock_pipeline import send_group_limit_alerts_for_items

    send_group_limit_alerts_for_items([])
    assert sends == []


def test_send_group_limit_alerts_send_error_does_not_propagate(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Этап 8: исключение из ветки отправки не пробрасывается наружу."""

    def boom(_html: str) -> None:
        raise RuntimeError("send")

    monkeypatch.setattr("services.stock_pipeline.list_all_limits", lambda: {"А": 5})
    monkeypatch.setattr(
        "services.stock_pipeline.find_group_limit_violations",
        lambda _items, _limits: [
            {
                "group": "А",
                "limit": 5,
                "positions": [{"name": "x", "quantity": 1}],
            },
        ],
    )
    monkeypatch.setattr(
        "services.stock_pipeline.format_group_limit_alerts_html",
        lambda _v: "<b>x</b>",
    )
    monkeypatch.setattr(
        "services.stock_pipeline.send_group_limit_alert_html",
        boom,
    )
    from services.stock_pipeline import send_group_limit_alerts_for_items

    send_group_limit_alerts_for_items(
        [{"top_category": "А", "name": "x", "quantity": 1}],
    )
